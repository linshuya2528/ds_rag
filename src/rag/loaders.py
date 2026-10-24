"""Load supported local course-material formats into page-level records."""

from __future__ import annotations

import hashlib
from pathlib import Path

from .schemas import Page

SUPPORTED_SUFFIXES = {".pdf", ".pptx", ".md", ".markdown", ".txt"}


def document_id_for(path: Path) -> str:
    """Return a stable content hash, independent of the file name."""
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()[:16]


def load_document(path: str | Path) -> list[Page]:
    file_path = Path(path)
    if not file_path.is_file():
        raise FileNotFoundError(f"Document does not exist: {file_path}")
    suffix = file_path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise ValueError(f"Unsupported document type {suffix!r}; supported: {supported}")
    document_id = document_id_for(file_path)
    if suffix == ".pdf":
        return _load_pdf(file_path, document_id)
    if suffix == ".pptx":
        return _load_pptx(file_path, document_id)
    return _load_text(file_path, document_id)


def _load_pdf(path: Path, document_id: str) -> list[Page]:
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError("Install pymupdf to ingest PDF files.") from exc
    pages: list[Page] = []
    with fitz.open(path) as document:
        for index, pdf_page in enumerate(document, start=1):
            text = pdf_page.get_text("text", sort=True).strip()
            if text:
                pages.append(Page(document_id, path.name, index, text))
    return pages


def _load_pptx(path: Path, document_id: str) -> list[Page]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("Install python-pptx to ingest PowerPoint files.") from exc
    presentation = Presentation(path)
    pages: list[Page] = []
    for index, slide in enumerate(presentation.slides, start=1):
        text = "\n".join(
            shape.text.strip()
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text.strip()
        )
        if text:
            pages.append(Page(document_id, path.name, index, text))
    return pages


def _load_text(path: Path, document_id: str) -> list[Page]:
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    return [Page(document_id, path.name, 1, text)] if text else []
